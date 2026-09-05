from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from pptx import Presentation

from models import Project
from services.slide_clone import duplicate_slide, remove_slide, set_bullet_paragraphs
from services.table_rows import set_data_row_count

TEMPLATE_PATH = Path(__file__).parent.parent.parent / "reference" / "CVChem-Mock-Updates.pptx"
EXPORT_DIR = Path(__file__).parent.parent / "data" / "exports"

TABLE_COLUMNS = ["_sno", "project_code", "_target_structure", "team_lead", "quantity", "start_date", "end_date", "on_time", "remarks"]


def _split_into_bullets(text: str | None) -> list[str]:
    if not text:
        return []
    parts = [s.strip() for s in re.split(r"(?<=[.])\s+", text.strip()) if s.strip()]
    return parts


def _fmt_date(value: dt.date | None) -> str:
    return value.strftime("%d-%m-%Y") if value else ""


def _fmt_on_time(value: bool | None) -> str:
    if value is None:
        return ""
    return "Yes" if value else "NO"


def _set_cell_text(table, row: int, col: int, text: str) -> None:
    table.cell(row, col).text = text


def _fill_master_table(table, projects: list[Project]) -> None:
    set_data_row_count(table, len(projects))
    for i, project in enumerate(projects):
        row = i + 1  # row 0 is the header
        _set_cell_text(table, row, 0, str(i + 1))
        _set_cell_text(table, row, 1, project.project_code)
        _set_cell_text(table, row, 2, project.target_molecule or "")
        _set_cell_text(table, row, 3, project.team_lead or "")
        _set_cell_text(table, row, 4, project.quantity or "")
        _set_cell_text(table, row, 5, _fmt_date(project.start_date))
        _set_cell_text(table, row, 6, _fmt_date(project.end_date))
        _set_cell_text(table, row, 7, _fmt_on_time(project.on_time))
        _set_cell_text(table, row, 8, project.remarks or "")


def _find_shape(slide, predicate):
    for shape in slide.shapes:
        if shape.has_text_frame and predicate(shape.text_frame.text.strip()):
            return shape
    return None


def _fill_project_slide(slide, project: Project) -> None:
    label_shape = _find_shape(slide, lambda t: t.startswith("Project:"))
    if label_shape is not None:
        tf = label_shape.text_frame
        if tf.paragraphs[0].runs:
            tf.paragraphs[0].runs[0].text = f"Project: {project.project_code}"
        else:
            tf.text = f"Project: {project.project_code}"

    overview_shape = _find_shape(slide, lambda t: t.startswith("Overview"))
    if overview_shape is not None:
        set_bullet_paragraphs(overview_shape.text_frame, _split_into_bullets(project.remarks))


def generate_rollup_pptx(projects: list[Project], output_name: str | None = None) -> Path:
    prs = Presentation(TEMPLATE_PATH)

    slide1 = prs.slides[0]
    date_shape = next((s for s in slide1.shapes if s.name.startswith("Date Placeholder")), None)
    if date_shape is not None and date_shape.text_frame.paragraphs[0].runs:
        date_shape.text_frame.paragraphs[0].runs[0].text = dt.date.today().strftime("%B %-d, %Y")

    slide2 = prs.slides[1]
    table_shape = next(s for s in slide2.shapes if s.has_table)
    _fill_master_table(table_shape.table, projects)

    prototype = prs.slides[2]
    new_slides = [duplicate_slide(prs, prototype) for _ in projects]

    original_project_slides = list(prs.slides)[2:11]  # the 9 original template project slides
    for old_slide in original_project_slides:
        remove_slide(prs, old_slide)

    for slide, project in zip(new_slides, projects):
        _fill_project_slide(slide, project)

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    output_name = output_name or f"rollup-{dt.date.today().isoformat()}.pptx"
    output_path = EXPORT_DIR / output_name
    prs.save(output_path)
    return output_path
