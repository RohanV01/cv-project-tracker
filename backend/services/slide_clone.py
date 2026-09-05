"""Low-level slide duplication for python-pptx.

python-pptx's public API has no way to duplicate an existing, fully-authored
slide (shapes, tables, images) with full fidelity — `Presentation.slides.add_slide(layout)`
only instantiates a layout's placeholders. This module is the one place in the
codebase that reaches into python-pptx internals to do real slide cloning; every
other module only does high-level text/table mutation on the results.
"""
from __future__ import annotations

import copy

from pptx.oxml.ns import qn
from pptx.presentation import Presentation
from pptx.slide import Slide
from pptx.text.text import _Paragraph
from pptx.text.text import TextFrame

RELATIONSHIP_TYPE_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"


def duplicate_slide(prs: Presentation, source_slide: Slide) -> Slide:
    """Append a deep copy of `source_slide` to `prs` and return the new Slide."""
    new_slide = prs.slides.add_slide(source_slide.slide_layout)

    # add_slide() pre-populates layout placeholders we don't want; we're about
    # to deep-copy the real shape tree from the source slide instead.
    for shape in list(new_slide.shapes):
        shape._element.getparent().remove(shape._element)

    old_to_new_rid: dict[str, str] = {}
    for shape in source_slide.shapes:
        new_shape_el = copy.deepcopy(shape._element)
        _remap_image_relationships(new_shape_el, source_slide, new_slide, old_to_new_rid)
        new_slide.shapes._spTree.append(new_shape_el)

    return new_slide


def _remap_image_relationships(shape_el, source_slide: Slide, new_slide: Slide, cache: dict[str, str]) -> None:
    for blip in shape_el.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        if rid not in cache:
            image_part = source_slide.part.related_part(rid)
            new_rid = new_slide.part.relate_to(image_part, RELATIONSHIP_TYPE_IMAGE)
            cache[rid] = new_rid
        blip.set(qn("r:embed"), cache[rid])


def set_bullet_paragraphs(text_frame: TextFrame, bullets: list[str]) -> None:
    """Keep text_frame.paragraphs[0] (the "Overview:" header) untouched, and
    replace every following paragraph with one bullet per string in `bullets`,
    cloning paragraph[1]'s XML as the formatting template so bullet styling
    (indent/bullet glyph from the slide master) is preserved."""
    paragraphs = text_frame.paragraphs
    if len(paragraphs) < 2:
        return  # prototype has no bullet paragraph to use as a template; leave as-is

    template_el = copy.deepcopy(paragraphs[1]._p)
    body = text_frame._txBody
    for p in paragraphs[1:]:
        body.remove(p._p)

    if not bullets:
        bullets = ["No updates recorded yet."]

    for bullet_text in bullets:
        new_p_el = copy.deepcopy(template_el)
        body.append(new_p_el)
        new_para = _Paragraph(new_p_el, text_frame)
        runs = new_para.runs
        if runs:
            runs[0].text = bullet_text
            for extra in runs[1:]:
                extra._r.getparent().remove(extra._r)
        else:
            new_para.add_run().text = bullet_text


def remove_slide(prs: Presentation, slide: Slide) -> None:
    """Remove `slide` from the presentation entirely (XML id list + part)."""
    xml_slides = prs.slides._sldIdLst
    slide_id_elements = list(xml_slides)
    for sld_id_el in slide_id_elements:
        if prs.part.related_part(sld_id_el.get(qn("r:id"))) is slide.part:
            xml_slides.remove(sld_id_el)
            prs.part.drop_rel(sld_id_el.get(qn("r:id")))
            break
