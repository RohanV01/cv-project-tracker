from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any

import openpyxl
from sqlalchemy.orm import Session

from models import IngestionRun, Project, ProjectHistory

# Column headers in the source workbook are not reliably descriptive (e.g. the
# "R&D" sheet's project-code column is literally labeled "CV Code " and its
# date columns are labeled "PO Date "/"PO Dispatch Date " even though they
# hold start/end dates) — so sheets are parsed by fixed column position,
# matched against the known layout, rather than by header text.

RD_FIELDS = ["_sno", "project_code", "quantity", "team_lead", "start_date", "end_date", "on_time", "remarks"]
PRODUCTION_FIELDS = [
    "_sno",
    "project_code",
    "old_or_new",
    "cas_no",
    "production_quantity_kg",
    "team_lead",
    "po_date",
    "po_dispatch_date",
    "dispatch_date",
    "delay_reason",
]

TRACKED_FIELDS = {
    "target_molecule",
    "team_lead",
    "quantity",
    "start_date",
    "end_date",
    "on_time",
    "remarks",
    "old_or_new",
    "cas_no",
    "production_quantity_kg",
    "po_date",
    "po_dispatch_date",
    "dispatch_date",
    "delay_reason",
}


def _coerce_date(value: Any) -> dt.date | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return None


def _coerce_on_time(value: Any) -> bool | None:
    if value is None or value == "":
        return None
    text = str(value).strip().upper()
    if text == "YES":
        return True
    if text == "NO":
        return False
    return None


def _coerce_str(value: Any) -> str | None:
    if value is None or value == "":
        return None
    return str(value).strip()


def _row_to_fields(row: tuple, field_names: list[str]) -> dict[str, Any]:
    raw = dict(zip(field_names, row))
    fields: dict[str, Any] = {"project_code": _coerce_str(raw.get("project_code"))}
    for name in field_names:
        if name in ("_sno", "project_code"):
            continue
        value = raw.get(name)
        if name in ("start_date", "end_date", "po_date", "po_dispatch_date", "dispatch_date"):
            fields[name] = _coerce_date(value)
        elif name == "on_time":
            fields[name] = _coerce_on_time(value)
        else:
            fields[name] = _coerce_str(value)
    return fields


def _diff(project: Project, fields: dict[str, Any]) -> dict[str, dict[str, Any]]:
    changed = {}
    for key, new_value in fields.items():
        old_value = getattr(project, key, None)
        old_cmp = old_value.isoformat() if isinstance(old_value, dt.date) else old_value
        new_cmp = new_value.isoformat() if isinstance(new_value, dt.date) else new_value
        if old_cmp != new_cmp:
            changed[key] = {"old": old_cmp, "new": new_cmp}
    return changed


def _snapshot(project: Project) -> dict[str, Any]:
    snap = {}
    for field in TRACKED_FIELDS:
        value = getattr(project, field, None)
        snap[field] = value.isoformat() if isinstance(value, dt.date) else value
    return snap


def upsert_project(
    session: Session, project_code: str, fields: dict[str, Any], source: str, source_file: str | None
) -> tuple[Project, bool, dict]:
    project = session.query(Project).filter_by(project_code=project_code).one_or_none()
    is_new = project is None
    if project is None:
        project = Project(project_code=project_code)
        session.add(project)
        session.flush()

    changed = {} if is_new else _diff(project, fields)
    for key, value in fields.items():
        setattr(project, key, value)
    session.flush()

    if is_new or changed:
        if is_new:
            changed = {
                k: {"old": None, "new": v.isoformat() if isinstance(v, dt.date) else v}
                for k, v in fields.items()
            }
        session.add(
            ProjectHistory(
                project_id=project.id,
                source=source,
                source_file=source_file,
                changed_fields=changed,
                snapshot=_snapshot(project),
            )
        )

    return project, is_new, changed


def _parse_sheet(ws, field_names: list[str]) -> list[dict[str, Any]]:
    rows = []
    for row in ws.iter_rows(min_row=2, max_col=len(field_names), values_only=True):
        fields = _row_to_fields(row, field_names)
        if not fields.get("project_code"):
            continue
        rows.append(fields)
    return rows


def ingest_workbook(path: Path | str, session: Session, filename: str | None = None) -> IngestionRun:
    filename = filename or Path(path).name
    wb = openpyxl.load_workbook(path, data_only=True)

    run = IngestionRun(filename=filename, sheet_names_found=wb.sheetnames, status="running")
    session.add(run)
    session.flush()

    created = updated = unchanged = skipped = processed = 0

    sheet_specs = []
    if "R&D" in wb.sheetnames:
        sheet_specs.append(("R&D", wb["R&D"], RD_FIELDS))
    if "Production" in wb.sheetnames:
        sheet_specs.append(("Production", wb["Production"], PRODUCTION_FIELDS))

    for _sheet_name, ws, field_names in sheet_specs:
        for fields in _parse_sheet(ws, field_names):
            processed += 1
            project_code = fields.pop("project_code")
            try:
                with session.begin_nested():
                    _project, is_new, changed = upsert_project(
                        session, project_code, fields, source="excel_import", source_file=filename
                    )
                if is_new:
                    created += 1
                elif changed:
                    updated += 1
                else:
                    unchanged += 1
            except Exception:
                skipped += 1

    run.rows_processed = processed
    run.rows_created = created
    run.rows_updated = updated
    run.rows_unchanged = unchanged
    run.rows_skipped_errors = skipped
    run.finished_at = dt.datetime.utcnow()
    run.status = "completed"
    session.commit()
    return run
