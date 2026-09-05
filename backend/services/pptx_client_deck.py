from __future__ import annotations

import datetime as dt
from pathlib import Path

from pptx import Presentation

from models import Project
from services.pptx_rollup import _fmt_date, _split_into_bullets
from services.slide_clone import set_bullet_paragraphs

TEMPLATE_PATH = Path(__file__).parent.parent.parent / "reference" / "CV-K23D-project-update-2026-06-18.pptx"
EXPORT_DIR = Path(__file__).parent.parent / "data" / "exports"


def generate_client_deck(project: Project, output_name: str | None = None) -> Path:
    prs = Presentation(TEMPLATE_PATH)

    slide1 = prs.slides[0]
    subtitle_shape = next((s for s in slide1.shapes if s.name.startswith("Subtitle")), None)
    if subtitle_shape is not None:
        subtitle_shape.text_frame.text = f"{project.project_code} Project Update\n{dt.date.today().strftime('%B %-d, %Y')}"
    date_shape = next((s for s in slide1.shapes if s.name.startswith("Date Placeholder")), None)
    if date_shape is not None and date_shape.text_frame.paragraphs[0].runs:
        date_shape.text_frame.paragraphs[0].runs[0].text = dt.date.today().strftime("%B %-d, %Y")

    slide2 = prs.slides[1]
    for shape in slide2.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip()
        if text.startswith("Overview"):
            set_bullet_paragraphs(shape.text_frame, _split_into_bullets(project.remarks))
        elif text.startswith("Target Qty"):
            lines = [
                f"Target Qty       :{project.quantity or '—'}",
                f"Assigned On  :{_fmt_date(project.start_date) or '—'}",
                f"Delivery date   :{_fmt_date(project.end_date) or '—'}",
            ]
            paragraphs = shape.text_frame.paragraphs
            for para, line in zip(paragraphs, lines):
                if para.runs:
                    para.runs[0].text = line
                    for extra in para.runs[1:]:
                        extra._r.getparent().remove(extra._r)

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    output_name = output_name or f"{project.project_code}-client-update-{dt.date.today().isoformat()}.pptx"
    output_path = EXPORT_DIR / output_name
    prs.save(output_path)
    return output_path
